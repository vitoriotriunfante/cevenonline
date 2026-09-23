import requests
import urllib3
import sqlite3
import openpyxl
import os
import json
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

TOKEN = 'IQ6vA2KkKhi1P03bC7WSMZNTw-i-xhJeDLHN3ga9rp4'
BASE_URL = 'https://ceven.drivetriunfante-locomotiva.com.br'
HEADERS = {'User-Agent': 'Mozilla/5.0'}

FILIAIS_MAP = {
    'ABC': 'abc1', 'API': 'api1', 'MCD': 'mcd1', 'TBE': 'tbe1',
    'TBL': 'tbl1', 'TCA': 'tca1', 'TCG': 'tcg1', 'TCV': 'tcv1',
    'TPA': 'tpa1', 'TPH': 'tph1', 'TSJ': 'tsj1'
}

db_path = os.path.join(os.path.dirname(__file__), 'pedidos_historico_ceven.db')
conn = sqlite3.connect(db_path)
cursor = conn.cursor()

print("="*80)
print("INICIANDO COLETA INTEGRAL (100% DOS VENDEDORES E TODAS AS CATEGORIAS/CNAES)")
print("="*80)

# 1. Carregar todos os vendedores da base
cursor.execute("""
    SELECT s.filial_sigla, s.rca_id, s.rca_nome, s.supervisor_nome, s.gerente_nome, s.segmento_nome
    FROM rca_segmentos s
    ORDER BY s.filial_sigla, s.rca_id
""")
vendedores = cursor.fetchall()
print(f"Total de vendedores na hierarquia oficial: {len(vendedores)}")

# 2. Coletar Roteiros de Hoje de TODOS os vendedores em paralelo
pdvs_hoje = []
rcas_com_rota = []

def fetch_roteiro(v):
    fil_sigla, rca_id, rca_nome, sup_nome, ger_nome, seg_nome = v
    fkey = FILIAIS_MAP.get(fil_sigla, fil_sigla.lower() + '1')
    url = f"{BASE_URL}/api/rca/roteiro-hoje?filial={fkey}&id={rca_id}"
    try:
        r = requests.get(url, headers=HEADERS, verify=False, timeout=12)
        if r.status_code == 200:
            data = r.json()
            if isinstance(data, list) and len(data) > 0:
                rows = []
                for item in data:
                    lat = item.get('latitude')
                    lon = item.get('longitude')
                    focos = item.get('focos') or []
                    focos_str = " | ".join([f.get('industria_foco', '') for f in focos if f.get('industria_foco')])
                    rows.append({
                        'Filial': fil_sigla,
                        'Gerente': ger_nome,
                        'Supervisor': sup_nome,
                        'RCA_Codigo': rca_id,
                        'RCA_Nome': rca_nome,
                        'Segmento': seg_nome,
                        'ID_Cliente': item.get('id_cliente') or item.get('id'),
                        'CNPJ': item.get('cnpj', ''),
                        'Nome_Fantasia': item.get('nome_cliente', ''),
                        'Razao_Social': item.get('razao_social', ''),
                        'Endereco': item.get('endereco', ''),
                        'Latitude': lat or '',
                        'Longitude': lon or '',
                        'Status_Visita': item.get('status', 'ABERTO'),
                        'Tempo_Visita': item.get('tempo_visita', ''),
                        'Checkout_Latitude': item.get('checkout_latitude', ''),
                        'Checkout_Longitude': item.get('checkout_longitude', ''),
                        'Data_Visita': item.get('data_visita', '2026-09-14'),
                        'Data_Ultima_Compra': (item.get('data_ultima_compra') or '')[:10],
                        'Focos_PEX': focos_str
                    })
                return (v, rows)
    except Exception:
        pass
    return (v, [])

print("\n--- COLETANDO ROTEIROS DE TODOS OS VENDEDORES (11 FILIAIS) ---")
t0 = time.time()
with ThreadPoolExecutor(max_workers=25) as executor:
    futures = [executor.submit(fetch_roteiro, v) for v in vendedores]
    for fut in as_completed(futures):
        v, rows = fut.result()
        if rows:
            pdvs_hoje.extend(rows)
            rcas_com_rota.append(v)

t_roteiros = time.time() - t0
print(f"Roteiros coletados em {t_roteiros:.2f}s:")
print(f" - Vendedores com rota ativa hoje: {len(rcas_com_rota)}")
print(f" - Total de PDVs na rota hoje: {len(pdvs_hoje)}")

# 3. Coletar Prospects de TODOS os vendedores com rota (SEM FILTRO DE CNAE)
print(f"\n--- CONSULTANDO RADAR DE PROSPECTS PARA TODOS OS {len(rcas_com_rota)} VENDEDORES COM ROTA ---")
prospects_coletados = []
cnpjs_prospects_vistos = set()

def fetch_prospects(v):
    fil_sigla, rca_id, rca_nome, sup_nome, ger_nome, seg_nome = v
    url = f"{BASE_URL}/api/ceven/prospeccao-roteiro?cod_rca={rca_id}&hoje=1&max=120&token={TOKEN}"
    try:
        r = requests.get(url, headers=HEADERS, verify=False, timeout=15)
        if r.status_code == 200:
            res_data = r.json()
            st = res_data.get('status')
            if st == 'pronto':
                return (v, res_data.get('prospects', []))
    except Exception:
        pass
    return (v, [])

t1 = time.time()
with ThreadPoolExecutor(max_workers=20) as executor:
    futures = [executor.submit(fetch_prospects, v) for v in rcas_com_rota]
    for fut in as_completed(futures):
        v, pro_list = fut.result()
        fil_sigla, rca_id, rca_nome, sup_nome, ger_nome, seg_nome = v
        for p in pro_list:
            cnpj_num = str(p.get('cnpj', '')).replace('.', '').replace('/', '').replace('-', '').strip()
            if cnpj_num and cnpj_num not in cnpjs_prospects_vistos:
                cnpjs_prospects_vistos.add(cnpj_num)
                cnpj_fmt = f"{cnpj_num[:2]}.{cnpj_num[2:5]}.{cnpj_num[5:8]}/{cnpj_num[8:12]}-{cnpj_num[12:14]}" if len(cnpj_num) == 14 else cnpj_num
                prospects_coletados.append({
                    '#': len(prospects_coletados) + 1,
                    'Filial': fil_sigla,
                    'Gerente': ger_nome,
                    'Supervisor': sup_nome,
                    'RCA Código': str(rca_id),
                    'Vendedor': rca_nome,
                    'CNPJ (Apenas Números)': cnpj_num,
                    'CNPJ (Formatado)': cnpj_fmt,
                    'Razão Social': p.get('razao') or p.get('nome') or '',
                    'Nome Fantasia': p.get('nome') or p.get('razao') or '',
                    'CNAE': p.get('cnae') or '',
                    'Descrição Atividade (CNAE)': p.get('cnae_desc') or '',
                    'Família CNAE': p.get('familia') or (str(p.get('cnae', ''))[:4]),
                    'Endereço': p.get('endereco') or '',
                    'Bairro': p.get('bairro') or '',
                    'Município': p.get('municipio') or '',
                    'UF': p.get('uf') or '',
                    'CEP': p.get('cep') or '',
                    'Telefone': p.get('telefone') or '',
                    'E-mail': p.get('email') or '',
                    'Distância Rota (km)': p.get('dist_km', 0),
                    'Grau Aderência': p.get('aderencia', 0),
                    'Motivo Radar CEVEN': p.get('motivo') or '',
                    'Status IE': p.get('ie_status') or 'nao_verificado',
                    'Fonte IE': p.get('ie_fonte') or 'pendente',
                    'Latitude': p.get('lat_aprox') or p.get('lat') or '',
                    'Longitude': p.get('lon_aprox') or p.get('lon') or '',
                    'Fonte Geo': p.get('geo_fonte') or 'radar_ceven'
                })

t_prospects = time.time() - t1
print(f"Prospects consultados em {t_prospects:.2f}s:")
print(f" - Total de prospects únicos capturados: {len(prospects_coletados)}")

# 4. Atualizar o Excel EXPANSAO_CADASTROS_E_PROSPECTS.xlsx (espelho legivel dos mesmos
# dados que vao pro SQLite abaixo -- nao pode ser bloqueante: se o arquivo nao existir
# nesta maquina (ex: runner do GitHub Actions, que nao tem esse arquivo local/manual),
# a atualizacao do SQLite (que e a fonte real usada pelo resto do pipeline) tem que
# continuar mesmo assim.
excel_path = os.path.join(os.path.dirname(__file__), 'EXPANSAO_CADASTROS_E_PROSPECTS.xlsx')
try:
    print(f"\n--- ATUALIZANDO PLANILHA EXCEL ({excel_path}) ---")
    wb = openpyxl.load_workbook(excel_path) if os.path.exists(excel_path) else openpyxl.Workbook()

    if 'PDVs_Roteiro_Hoje_Geolocalizados' in wb.sheetnames:
        del wb['PDVs_Roteiro_Hoje_Geolocalizados']

    ws_pdvs = wb.create_sheet('PDVs_Roteiro_Hoje_Geolocalizados')
    if pdvs_hoje:
        headers_pdvs = list(pdvs_hoje[0].keys())
        ws_pdvs.append(headers_pdvs)
        for row in pdvs_hoje:
            ws_pdvs.append([row[h] for h in headers_pdvs])
    print(f" - Aba 'PDVs_Roteiro_Hoje_Geolocalizados' regravada com {len(pdvs_hoje)} PDVs!")

    if 'Prospects_Mapa_Radar' in wb.sheetnames:
        del wb['Prospects_Mapa_Radar']

    ws_prosp = wb.create_sheet('Prospects_Mapa_Radar')
    if prospects_coletados:
        headers_prosp = list(prospects_coletados[0].keys())
        ws_prosp.append(headers_prosp)
        for row in prospects_coletados:
            ws_prosp.append([row[h] for h in headers_prosp])
    print(f" - Aba 'Prospects_Mapa_Radar' regravada com {len(prospects_coletados)} prospects de TODOS os CNAEs!")

    if 'Sheet' in wb.sheetnames and wb['Sheet'].max_row == 1 and wb['Sheet'].max_column == 1:
        del wb['Sheet']

    wb.save(excel_path)
    print(f"Planilha {excel_path} salva com sucesso!")
except Exception as e:
    print(f"⚠️  Não foi possível atualizar a planilha Excel ({e}) — seguindo para o SQLite normalmente.")

# 5. Atualizar SQLite
print("\n--- ATUALIZANDO BASE SQLITE ---")
cursor.execute("DROP TABLE IF EXISTS pdvs_roteiro_hoje_gps")
cursor.execute("""
    CREATE TABLE pdvs_roteiro_hoje_gps (
        filial TEXT, gerente TEXT, supervisor TEXT, rca_codigo TEXT, rca_nome TEXT,
        segmento TEXT, id_cliente TEXT, cnpj TEXT, nome_fantasia TEXT, razao_social TEXT,
        endereco TEXT, latitude REAL, longitude REAL, status_visita TEXT, tempo_visita TEXT,
        checkout_latitude REAL, checkout_longitude REAL, data_visita TEXT,
        data_ultima_compra TEXT, focos_pex TEXT
    )
""")
for p in pdvs_hoje:
    cursor.execute("""
        INSERT INTO pdvs_roteiro_hoje_gps VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
    """, (
        p.get('Filial'), p.get('Gerente'), p.get('Supervisor'), str(p.get('RCA_Codigo')), p.get('RCA_Nome'),
        p.get('Segmento'), str(p.get('ID_Cliente')), p.get('CNPJ'), p.get('Nome_Fantasia'), p.get('Razao_Social'),
        p.get('Endereco'), p.get('Latitude') or None, p.get('Longitude') or None, p.get('Status_Visita'),
        p.get('Tempo_Visita'), p.get('Checkout_Latitude') or None, p.get('Checkout_Longitude') or None,
        p.get('Data_Visita'), p.get('Data_Ultima_Compra'), p.get('Focos_PEX')
    ))

cursor.execute("DROP TABLE IF EXISTS prospects_mapa_radar")
cursor.execute("""
    CREATE TABLE prospects_mapa_radar (
        filial TEXT, gerente TEXT, supervisor TEXT, rca_codigo TEXT, vendedor TEXT,
        cnpj TEXT, cnpj_formatado TEXT, razao_social TEXT, nome_fantasia TEXT,
        cnae TEXT, cnae_desc TEXT, familia_cnae TEXT, endereco TEXT, bairro TEXT,
        municipio TEXT, uf TEXT, cep TEXT, telefone TEXT, email TEXT,
        distancia_km REAL, aderencia REAL, motivo TEXT, status_ie TEXT,
        fonte_ie TEXT, latitude REAL, longitude REAL, fonte_geo TEXT
    )
""")
for pr in prospects_coletados:
    cursor.execute("""
        INSERT INTO prospects_mapa_radar VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
    """, (
        pr.get('Filial'), pr.get('Gerente'), pr.get('Supervisor'), str(pr.get('RCA Código')), pr.get('Vendedor'),
        pr.get('CNPJ (Apenas Números)'), pr.get('CNPJ (Formatado)'), pr.get('Razão Social'), pr.get('Nome Fantasia'),
        pr.get('CNAE'), pr.get('Descrição Atividade (CNAE)'), pr.get('Família CNAE'), pr.get('Endereço'),
        pr.get('Bairro'), pr.get('Município'), pr.get('UF'), pr.get('CEP'), pr.get('Telefone'),
        pr.get('E-mail'), pr.get('Distância Rota (km)'), pr.get('Grau Aderência'), pr.get('Motivo Radar CEVEN'),
        pr.get('Status IE'), pr.get('Fonte IE'), pr.get('Latitude') or None, pr.get('Longitude') or None,
        pr.get('Fonte Geo')
    ))

conn.commit()
print("Base SQLite atualizada com sucesso!")

# Salvar backups JSON
with open(os.path.join(os.path.dirname(__file__), 'PDVS_ROTEIRO_HOJE_GEOLOCALIZADOS.json'), 'w', encoding='utf-8') as f:
    json.dump(pdvs_hoje, f, ensure_ascii=False, indent=2)

with open(os.path.join(os.path.dirname(__file__), 'PROSPECTS_RADAR_11_FILIAIS_HOJE.json'), 'w', encoding='utf-8') as f:
    json.dump(prospects_coletados, f, ensure_ascii=False, indent=2)

print("\nArquivos JSON de backup atualizados.")
print("="*80)
print("COLETA INTEGRAL CONCLUÍDA COM SUCESSO!")
print("="*80)
