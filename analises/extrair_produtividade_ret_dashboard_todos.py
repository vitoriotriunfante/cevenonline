import requests
import urllib3
import sqlite3
import openpyxl
import os
import json
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

urllib3.disable_warnings()

headers = {'User-Agent': 'Mozilla/5.0'}
base_url = 'https://ceven.drivetriunfante-locomotiva.com.br'

db_path = os.path.join(os.path.dirname(__file__), 'pedidos_historico_ceven.db')
conn = sqlite3.connect(db_path)
cursor = conn.cursor()

print("="*80)
print("EXTRAINDO DADOS OPERACIONAIS AO VIVO (PRODUTIVIDADE, RET E DASHBOARD FINANCEIRO)")
print("PARA TODOS OS VENDEDORES EM CAMPO (11 FILIAIS)")
print("="*80)

# Carregar os vendedores com rota hoje
cursor.execute("""
    SELECT DISTINCT filial, gerente, supervisor, rca_codigo, rca_nome, segmento
    FROM pdvs_roteiro_hoje_gps
    ORDER BY filial, rca_codigo
""")
vendedores = cursor.fetchall()
print(f"Total de vendedores com rota ativa hoje para consultar: {len(vendedores)}")

filiais_map = {
    'ABC': 'abc1', 'API': 'api1', 'MCD': 'mcd1', 'TBE': 'tbe1',
    'TBL': 'tbl1', 'TCA': 'tca1', 'TCG': 'tcg1', 'TCV': 'tcv1',
    'TPA': 'tpa1', 'TPH': 'tph1', 'TSJ': 'tsj1'
}

produtividade_lista = []
dashboard_lista = []
ret_hoje_lista = []

def fetch_vendedor_data(v):
    fil_sigla, ger_nome, sup_nome, rca_id, rca_nome, seg_nome = v
    fkey = filiais_map.get(fil_sigla, fil_sigla.lower() + '1')
    
    res = {'prod': None, 'dash': None, 'ret': None}
    
    # 1. Produtividade
    try:
        r1 = requests.get(f"{base_url}/api/rca/produtividade?filial={fkey}&id={rca_id}", headers=headers, verify=False, timeout=10)
        if r1.status_code == 200:
            res['prod'] = r1.json()
    except Exception:
        pass
        
    # 2. Dashboard Financeiro
    try:
        r2 = requests.get(f"{base_url}/api/rca/dashboard?filial={fkey}&id={rca_id}", headers=headers, verify=False, timeout=10)
        if r2.status_code == 200:
            res['dash'] = r2.json()
    except Exception:
        pass
        
    # 3. RET Hoje (Execução de Campo)
    try:
        r3 = requests.get(f"{base_url}/api/rca/ret-hoje?filial={fkey}&id={rca_id}", headers=headers, verify=False, timeout=10)
        if r3.status_code == 200:
            res['ret'] = r3.json()
    except Exception:
        pass
        
    return (v, res)

t0 = time.time()
print("\nIniciando requisições paralelas (25 workers)...")

with ThreadPoolExecutor(max_workers=25) as executor:
    futures = [executor.submit(fetch_vendedor_data, v) for v in vendedores]
    for fut in as_completed(futures):
        v, res = fut.result()
        fil_sigla, ger_nome, sup_nome, rca_id, rca_nome, seg_nome = v
        
        # Parse Produtividade
        prod = res.get('prod')
        if prod and isinstance(prod, dict):
            dia = prod.get('dia') or {}
            mes = prod.get('mes') or {}
            mix = prod.get('mix_mes') or {}
            produtividade_lista.append({
                'Filial': fil_sigla,
                'Gerente': ger_nome,
                'Supervisor': sup_nome,
                'RCA_Codigo': rca_id,
                'RCA_Nome': rca_nome,
                'Segmento': seg_nome,
                'Dia_Faturamento': dia.get('faturamento', 0),
                'Dia_Positivacao': dia.get('positivacao', 0),
                'Dia_Digitado_Pedido': dia.get('dig_pedido', 0),
                'Dia_Visitas_Programadas': dia.get('total_programado', 0),
                'Dia_Visitas_Rota': dia.get('visitas_na_rota', 0),
                'Dia_Visitas_Com_Venda': dia.get('visitas_com_venda', 0),
                'Dia_Eficacia_Pct': dia.get('eficacia_pct', 0),
                'Mes_Visitas_Planejadas': mes.get('visit_plan', 0),
                'Mes_Visitas_Realizadas': mes.get('visit_real', 0),
                'Mes_Pedidos_Rota': mes.get('ped_rota_mes', 0),
                'Mes_Eficiencia_Visitas_Pct': mes.get('eficiencia_pct', 0),
                'Mes_SKUs_Distintos': mix.get('skus_distintos', 0),
                'Mes_Positivacao_Media': mix.get('positivacao_media', 0)
            })
            
        # Parse Dashboard
        dash = res.get('dash')
        if dash and isinstance(dash, dict):
            fin = dash.get('financeiro') or {}
            pos = dash.get('positivacao') or {}
            dashboard_lista.append({
                'Filial': fil_sigla,
                'Gerente': ger_nome,
                'Supervisor': sup_nome,
                'RCA_Codigo': rca_id,
                'RCA_Nome': rca_nome,
                'Segmento': seg_nome,
                'Meta_Faturamento': fin.get('meta', 0),
                'Faturado_Liquido': fin.get('faturado', 0),
                'Faturado_Bruto': fin.get('faturado_bruto', 0),
                'Devolucao_Valor': fin.get('devolucao', 0),
                'Pendente_Faturamento': fin.get('pendente', 0),
                'Total_Com_Pendente': fin.get('total_com_pendente', 0),
                'Atingimento_Financ_Pct': fin.get('atingimento_pct', 0),
                'Faltante_Meta_Financ': fin.get('faltante', 0),
                'Meta_Positivacao': pos.get('meta', 0),
                'Positivacao_Realizada': pos.get('realizada', 0),
                'Atingimento_Posit_Pct': pos.get('atingimento_pct', 0),
                'Carteira_Total_Clientes': pos.get('carteira_total', 0),
                'Clientes_Inativos': pos.get('inativos', 0)
            })
            
        # Parse RET Hoje
        ret = res.get('ret')
        if ret and isinstance(ret, dict):
            ov = ret.get('overview') or {}
            ret_hoje_lista.append({
                'Filial': fil_sigla,
                'Gerente': ger_nome,
                'Supervisor': sup_nome,
                'RCA_Codigo': rca_id,
                'RCA_Nome': rca_nome,
                'Segmento': seg_nome,
                'PDVs_Visitados_Ate_Agora': ov.get('pdvsVisitados', 0),
                'Duracao_Total_Minutos': ov.get('duracaoTotalMin', 0),
                'Primeiro_Checkin': ov.get('primeiroCheckin', ''),
                'Ultimo_Checkout': ov.get('ultimoCheckout', ''),
                'Total_Visitas_Logadas': len(ret.get('visitas', []))
            })

t_fim = time.time() - t0
print(f"\nExtração concluída em {t_fim:.2f}s!")
print(f" - Produtividade extraída: {len(produtividade_lista)} vendedores")
print(f" - Dashboards financeiros extraídos: {len(dashboard_lista)} vendedores")
print(f" - Status RET campo extraídos: {len(ret_hoje_lista)} vendedores")

# 1. Salvar no Banco SQLite
print("\n--- SALVANDO NO BANCO SQLITE pedidos_historico_ceven.db ---")

cursor.execute("DROP TABLE IF EXISTS rca_produtividade_live")
cursor.execute("""
    CREATE TABLE rca_produtividade_live (
        filial TEXT, gerente TEXT, supervisor TEXT, rca_codigo TEXT, rca_nome TEXT, segmento TEXT,
        dia_faturamento REAL, dia_positivacao INTEGER, dia_digitado_pedido REAL,
        dia_visitas_programadas INTEGER, dia_visitas_rota INTEGER, dia_visitas_com_venda INTEGER,
        dia_eficacia_pct REAL, mes_visitas_planejadas INTEGER, mes_visitas_realizadas INTEGER,
        mes_pedidos_rota INTEGER, mes_eficiencia_visitas_pct REAL, mes_skus_distintos INTEGER,
        mes_positivacao_media REAL
    )
""")
for p in produtividade_lista:
    cursor.execute("""
        INSERT INTO rca_produtividade_live VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
    """, tuple(p.values()))

cursor.execute("DROP TABLE IF EXISTS rca_dashboard_financeiro_live")
cursor.execute("""
    CREATE TABLE rca_dashboard_financeiro_live (
        filial TEXT, gerente TEXT, supervisor TEXT, rca_codigo TEXT, rca_nome TEXT, segmento TEXT,
        meta_faturamento REAL, faturado_liquido REAL, faturado_bruto REAL, devolucao_valor REAL,
        pendente_faturamento REAL, total_com_pendente REAL, atingimento_financ_pct REAL,
        faltante_meta_financ REAL, meta_positivacao INTEGER, positivacao_realizada INTEGER,
        atingimento_posit_pct REAL, carteira_total_clientes INTEGER, clientes_inativos INTEGER
    )
""")
for d in dashboard_lista:
    cursor.execute("""
        INSERT INTO rca_dashboard_financeiro_live VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
    """, tuple(d.values()))

cursor.execute("DROP TABLE IF EXISTS rca_ret_execucao_hoje")
cursor.execute("""
    CREATE TABLE rca_ret_execucao_hoje (
        filial TEXT, gerente TEXT, supervisor TEXT, rca_codigo TEXT, rca_nome TEXT, segmento TEXT,
        pdvs_visitados_ate_agora INTEGER, duracao_total_minutos REAL, primeiro_checkin TEXT,
        ultimo_checkout TEXT, total_visitas_logadas INTEGER
    )
""")
for r in ret_hoje_lista:
    cursor.execute("""
        INSERT INTO rca_ret_execucao_hoje VALUES (?,?,?,?,?,?,?,?,?,?,?)
    """, tuple(r.values()))

conn.commit()
print("Tabelas SQLite atualizadas com sucesso!")

# 2. Gerar Planilha Excel Consolidada
excel_out = os.path.join(os.path.dirname(__file__), 'PRODUTIVIDADE_E_RET_AO_VIVO_14_09.xlsx')
print(f"\n--- GERANDO PLANILHA EXCEL ({excel_out}) ---")
wb = openpyxl.Workbook()
wb.remove(wb.active)

# Aba 1: Dashboard Financeiro Live
ws_dash = wb.create_sheet('Dashboard_Financeiro_Live')
if dashboard_lista:
    h_dash = list(dashboard_lista[0].keys())
    ws_dash.append(h_dash)
    for row in dashboard_lista:
        ws_dash.append([row[h] for h in h_dash])

# Aba 2: Produtividade Rota e Mês
ws_prod = wb.create_sheet('Produtividade_Diaria_e_Mes')
if produtividade_lista:
    h_prod = list(produtividade_lista[0].keys())
    ws_prod.append(h_prod)
    for row in produtividade_lista:
        ws_prod.append([row[h] for h in h_prod])

# Aba 3: RET Campo Hoje
ws_ret = wb.create_sheet('RET_Execucao_Campo_Hoje')
if ret_hoje_lista:
    h_ret = list(ret_hoje_lista[0].keys())
    ws_ret.append(h_ret)
    for row in ret_hoje_lista:
        ws_ret.append([row[h] for h in h_ret])

wb.save(excel_out)
print(f"Planilha Excel {excel_out} gerada com sucesso!")

# 3. Salvar backups JSON
with open(os.path.join(os.path.dirname(__file__), 'PRODUTIVIDADE_RCA_LIVE.json'), 'w', encoding='utf-8') as f:
    json.dump(produtividade_lista, f, ensure_ascii=False, indent=2)

with open(os.path.join(os.path.dirname(__file__), 'DASHBOARD_FINANCEIRO_LIVE.json'), 'w', encoding='utf-8') as f:
    json.dump(dashboard_lista, f, ensure_ascii=False, indent=2)

with open(os.path.join(os.path.dirname(__file__), 'RET_EXECUCAO_HOJE.json'), 'w', encoding='utf-8') as f:
    json.dump(ret_hoje_lista, f, ensure_ascii=False, indent=2)

print("Arquivos JSON de backup salvos.")
print("="*80)
print("CARGA OPERACIONAL 100% CONCLUÍDA COM SUCESSO!")
print("="*80)
